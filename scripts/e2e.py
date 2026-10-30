import asyncio
import hashlib
import json
import os
from pathlib import Path
import platform
import shlex
import subprocess
import sys
import time
import urllib.error
import urllib.request
import zipfile

from docx import Document
from mcp import ClientSession, StdioServerParameters
from mcp.client.stdio import stdio_client
from openpyxl import load_workbook
from PIL import Image
import psutil
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parents[1]
WORK = ROOT / ".work/e2e"
VARIANT = os.environ.get("DC_TEST_VARIANT","desktop")
EVIDENCE = ROOT / ".work/evidence" / VARIANT
BUNDLE = ROOT / ".work/bundle"
EXE = Path(os.environ.get("DC_TEST_EXECUTABLE", BUNDLE / ("desktop-commander.exe" if os.name == "nt" else "desktop-commander")))
REPORT = {"platform": platform.platform(), "variant":VARIANT, "commit":os.environ.get("GITHUB_SHA"), "passed":False, "cases": [], "artifacts": {}}


def record(name, details=None):
    REPORT["cases"].append({"name": name, "passed": True, "details": details})
    print(name, flush=True)


def structured(response):
    if response.isError:
        raise AssertionError(response.model_dump(by_alias=True))
    data = response.model_dump(by_alias=True).get("structuredContent")
    if data is not None:
        return data
    return json.loads(response.content[0].text)


def shell_command(arguments):
    if os.name == "nt":
        return "& " + " ".join("'"+argument.replace("'", "''")+"'" for argument in arguments)
    return shlex.join(arguments)


async def mcp_flow():
    config = WORK / "stdio-config"
    parameters = StdioServerParameters(command=str(EXE), args=["--config-dir", str(config), "stdio"])
    with (EVIDENCE / "mcp-stderr.log").open("w") as errors:
        async with stdio_client(parameters, errlog=errors) as (reader, writer):
            async with ClientSession(reader, writer) as client:
                await client.initialize()
                tools = await client.list_tools()
                expected = json.loads((ROOT / "reference/tools.json").read_text(encoding="utf-8"))
                assert {tool.name for tool in tools.tools} == {tool["name"] for tool in expected}
                record("MCP tool inventory", len(tools.tools))
                for resource in (await client.list_resources()).resources:
                    content = await client.read_resource(resource.uri)
                    assert "<html" in content.contents[0].text.lower()
                record("MCP HTML resources")

                async def call(name, **args):
                    response = await client.call_tool(name, args)
                    if response.isError:
                        raise AssertionError(f"{name}: {response.model_dump(by_alias=True)}")
                    return response

                path = WORK / "files"
                await call("create_directory", path=str(path))
                source = path / "文本.md"
                await call("write_file", path=str(source), content="# 中文测试\nalpha\nbeta\n")
                response = await call("read_file", path=str(source), offset=1, length=1)
                assert response.content[0].text == "alpha"
                await call("edit_block", file_path=str(source), old_string="alpha", new_string="updated")
                assert "updated" in source.read_text(encoding="utf-8")
                await call("write_file", path=str(source), content="append\n", mode="append")
                await call("read_multiple_files", paths=[str(source)])
                await call("get_file_info", path=str(source))
                await call("list_directory", path=str(path), depth=2)
                destination = path / "moved.md"
                await call("move_file", source=str(source), destination=str(destination))
                assert destination.exists() and not source.exists()
                failure = await client.call_tool("edit_block", {"file_path":str(destination), "old_string":"absent", "new_string":"x"})
                assert failure.isError and "absent" not in destination.read_text(encoding="utf-8")
                record("Filesystem creation, paging, append, edit, move, failure")

                search = structured(await call("start_search", path=str(path), pattern="updated", searchType="content", timeout_ms=1000))
                search_id = search["sessionId"]
                for _ in range(30):
                    search = structured(await call("get_more_search_results", sessionId=search_id))
                    if search["isComplete"]:
                        break
                    await asyncio.sleep(0.1)
                assert any("updated" in entry["text"] for entry in search["results"])
                await call("list_searches")
                await call("stop_search", sessionId=search_id)
                record("Real ripgrep search and pagination")

                command = shell_command([sys.executable,"-u","-c","print('terminal-e2e')"])
                terminal = structured(await call("start_process", command=command, timeout_ms=2000,verbose_timing=True))
                assert terminal["timingInfo"]["totalDurationMs"] >= 0
                output = terminal["output"]
                for _ in range(30):
                    if terminal["exitCode"] is not None:
                        break
                    terminal = structured(await call("read_process_output",pid=terminal["pid"],timeout_ms=1000))
                    output += terminal["output"]
                assert "terminal-e2e" in output, terminal
                measured=structured(await call("read_process_output",pid=terminal["pid"],offset=-1,verbose_timing=True))
                assert measured["timingInfo"]["outputEvents"]
                interactive = shell_command([sys.executable,"-u","-c","print(input(), flush=True)"])
                terminal = structured(await call("start_process", command=interactive, timeout_ms=100))
                response = structured(await call("interact_with_process", pid=terminal["pid"], input="interactive-e2e", timeout_ms=2000))
                assert "interactive-e2e" in response["output"]
                await call("read_process_output", pid=terminal["pid"], offset=-1, length=1)
                await call("list_sessions")
                await call("force_terminate", pid=terminal["pid"])
                await call("list_processes")
                record("Real shell execution, stdin interaction and termination")
                streaming_code = "import sys,time;sys.stdout.write('x'*2097152);sys.stdout.flush();time.sleep(1);print('stream-tail',flush=True)"
                streaming = shell_command([sys.executable,"-u","-c",streaming_code])
                first = structured(await call("start_process", command=streaming, timeout_ms=500))
                tail = structured(await call("read_process_output",pid=first["pid"],timeout_ms=2000))
                assert "stream-tail" in first["output"]+tail["output"]
                assert tail["discardedBytes"] >= 1024*1024
                record("Bounded terminal output and incremental output after eviction")

                unicode_code = "import sys,time;sys.stdout.buffer.write(bytes([228]));sys.stdout.flush();time.sleep(0.8);sys.stdout.buffer.write(bytes([184,173]));sys.stdout.flush();time.sleep(0.8);sys.stdout.buffer.write(bytes([228]));sys.stdout.flush()"
                terminal = structured(await call("start_process",command=shell_command([sys.executable,"-u","-c",unicode_code]),timeout_ms=100))
                output = terminal["output"]
                assert "\ufffd" not in output
                for _ in range(10):
                    terminal = structured(await call("read_process_output",pid=terminal["pid"],timeout_ms=500))
                    output += terminal["output"]
                    if terminal["exitCode"] is not None:
                        break
                assert terminal["exitCode"] == 0 and output == "中\ufffd", (terminal,output)
                assert structured(await call("read_process_output",pid=terminal["pid"]))["output"] == ""
                record("UTF-8 streaming preserves split characters and consumes incomplete EOF")

                docx = path / "report.docx"
                await call("write_file", path=str(docx), content="Document E2E\nOriginal paragraph")
                assert "Original paragraph" in "\n".join(p.text for p in Document(docx).paragraphs)
                original = zipfile.ZipFile(docx)
                untouched = {name: original.read(name) for name in original.namelist() if name != "word/document.xml"}
                original.close()
                docx_content = structured(await call("read_file", path=str(docx)))
                assert docx_content["fileType"] == "docx"
                assert docx_content["mimeType"] == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                await call("edit_block", file_path=str(docx), old_string="Original paragraph", new_string="Changed paragraph")
                assert "Changed paragraph" in "\n".join(p.text for p in Document(docx).paragraphs)
                with zipfile.ZipFile(docx) as archive:
                    assert all(archive.read(name) == content for name, content in untouched.items())
                await call("get_file_info", path=str(docx))
                record("DOCX real round trip and untouched part preservation")

                xlsx = path / "data.xlsx"
                await call("write_file", path=str(xlsx), content=json.dumps([["name","value","enabled"],["sample",12,True]]))
                await call("edit_block", file_path=str(xlsx), range="Sheet1!B2", content=[[42]])
                await call("write_file", path=str(xlsx), content=json.dumps([["extra",23,False]]), mode="append")
                await call("read_file", path=str(xlsx), range="A1:C3")
                book = load_workbook(xlsx)
                assert book["Sheet1"]["B2"].value == 42
                assert book["Sheet1"]["C2"].value is True
                assert book["Sheet1"]["A3"].value == "extra"
                await call("get_file_info", path=str(xlsx))
                record("XLSX real write, range edit, append and types")

                pdf = path / "report.pdf"
                png = path / "image.png"
                Image.new("RGB",(120,80),(90,100,180)).save(png)
                await call("write_pdf", path=str(pdf), content="# PDF E2E\n\nNative typography test.\n\n中文排版验证。\n\n![image](image.png)", options={"margin":"18mm"})
                reader = PdfReader(pdf)
                assert "Native typography" in reader.pages[0].extract_text()
                assert len(reader.pages[0].images) > 0
                extracted = await call("read_file", path=str(pdf), offset=0, length=1)
                assert any(content.type == "image" for content in extracted.content)
                assert structured(extracted)["fileType"] == "pdf"
                assert structured(extracted)["mimeType"] == "application/pdf"
                await call("write_pdf", path=str(pdf), content=[{"type":"insert","pageIndex":1,"markdown":"Inserted page"}])
                assert len(PdfReader(pdf).pages) == 2
                await call("write_pdf", path=str(pdf), content=[{"type":"delete","pageIndexes":[0]}])
                assert len(PdfReader(pdf).pages) == 1
                assert "Inserted" in PdfReader(pdf).pages[0].extract_text()
                await call("get_file_info", path=str(pdf))
                record("Native PDF creation, extraction, insertion and deletion")

                response = await call("read_file", path=str(png))
                assert response.content[0].type == "image"
                record("Image decoding and MCP image response")

                config_data = structured(await call("get_config"))
                await call("set_config_value", key="fileReadLineLimit", value=300)
                assert structured(await call("get_config"))["config"]["fileReadLineLimit"] == 300
                await call("get_recent_tool_calls", maxResults=10)
                await call("get_usage_stats")
                await call("get_prompts", action="get_prompt", promptId="onb2_01")
                await call("track_ui_event", event="e2e_completed")
                record("Configuration persistence, history, prompts")
                for item in path.iterdir():
                    if item.is_file():
                        copied = EVIDENCE / item.name
                        copied.write_bytes(item.read_bytes())
                        REPORT["artifacts"][item.name] = hashlib.sha256(copied.read_bytes()).hexdigest()


def control_flow():
    config = WORK / "control-config"
    config.mkdir(parents=True, exist_ok=True)
    (config / "settings.json").write_text(json.dumps({"autoConnect":False}), encoding="utf-8")
    with (EVIDENCE / "control-stderr.log").open("w") as errors:
        child = subprocess.Popen([str(EXE),"--config-dir",str(config),"remote","--headless"], stderr=errors)
        try:
            descriptor = config / "control.json"
            for _ in range(100):
                if child.poll() is not None:
                    raise AssertionError("Control process exited")
                if descriptor.exists():
                    break
                time.sleep(0.1)
            connection = json.loads(descriptor.read_text(encoding="utf-8"))
            opener = urllib.request.build_opener(urllib.request.ProxyHandler({}))

            def request(path, method="GET", data=None, token=True):
                headers = {"Authorization": f"Bearer {connection['token']}"} if token else {}
                if data is not None:
                    headers["Content-Type"] = "application/json"
                request = urllib.request.Request(connection["endpoint"]+path, method=method,
                    headers=headers, data=json.dumps(data).encode() if data is not None else None)
                with opener.open(request, timeout=10) as response:
                    return json.load(response)

            unauthorized = False
            try:
                request("/status",token=False)
            except urllib.error.HTTPError as error:
                unauthorized = error.code == 401
            assert unauthorized
            status = request("/status")
            assert status["status"]["state"] == "stopped"
            duplicate = subprocess.run([str(EXE),"--config-dir",str(config),"remote","--headless"],capture_output=True,timeout=10)
            assert duplicate.returncode != 0
            assert request("/status")["pid"] == child.pid
            record("Single instance rejects a duplicate headless process")
            settings = status["settings"]
            settings["deviceName"] = "E2E native device"
            request("/settings","PUT",settings)
            assert request("/status")["settings"]["deviceName"] == "E2E native device"
            exported = Path(request("/export","POST")["path"])
            assert "access_token" not in exported.read_text(encoding="utf-8")
            record("Real authenticated control API and persisted settings")
            process = psutil.Process(child.pid)
            samples = []
            for _ in range(10):
                family = [process] + process.children(recursive=True)
                memory = sum(p.memory_full_info().pss if sys.platform.startswith("linux") else p.memory_info().rss for p in family)
                samples.append(memory)
                time.sleep(0.2)
            REPORT["idleStoppedMemory"] = {"metric":"PSS" if sys.platform.startswith("linux") else "RSS",
                "samplesBytes":samples,"maxMiB":max(samples)/1024**2,"connected":False}
            request("/quit","POST")
            child.wait(timeout=10)
            assert child.returncode == 0
            assert not descriptor.exists()
            record("Clean shutdown and descriptor cleanup")
            if os.name != "nt":
                restarted = subprocess.Popen([str(EXE),"--config-dir",str(config),"remote","--headless"],stderr=errors)
                try:
                    for _ in range(100):
                        if descriptor.exists():
                            break
                        time.sleep(0.1)
                    assert descriptor.exists()
                    restarted.terminate()
                    restarted.wait(timeout=10)
                    assert restarted.returncode == 0
                    assert not descriptor.exists()
                    record("SIGTERM performs graceful cleanup")
                finally:
                    if restarted.poll() is None:
                        restarted.kill()
                        restarted.wait()
        finally:
            if child.poll() is None:
                child.kill()
                child.wait()


def main():
    WORK.mkdir(parents=True, exist_ok=True)
    EVIDENCE.mkdir(parents=True, exist_ok=True)
    try:
        asyncio.run(mcp_flow())
        control_flow()
        REPORT["passed"] = True
    finally:
        (EVIDENCE / "e2e.json").write_text(json.dumps(REPORT,ensure_ascii=False,indent=2),encoding="utf-8")


if __name__ == "__main__":
    main()
